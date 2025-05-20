import streamlit as st
from docx import Document
from io import BytesIO
from datetime import datetime
import os
import subprocess

# Funzione per convertire DOCX in PDF usando libreoffice (compatibile Linux)
def convert_to_pdf(input_path, output_path):
    subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(output_path), input_path])

st.title("Generatore Contratto di Procacciamento")

nome = st.text_input("Nome completo del Procacciatore")
sede = st.text_input("Sede del Procacciatore")
via = st.text_input("Via")
piva = st.text_input("Partita IVA")
cf = st.text_input("Codice Fiscale")
prov_gas = st.text_input("Provvigione Gas (€/mc)")
prov_luce = st.text_input("Provvigione Luce (€/pod)")
data = st.date_input("Data di sottoscrizione", value=datetime.today())

if st.button("Genera PDF"):
    doc = Document("contratto_template.docx")
    
    placeholders = {
        "{{NOME}}": nome,
        "{{SEDE}}": sede,
        "{{VIA}}": via,
        "{{PIVA}}": piva,
        "{{CF}}": cf,
        "{{PROV_GAS}}": prov_gas,
        "{{PROV_LUCE}}": prov_luce,
        "{{DATA}}": data.strftime("%d/%m/%Y"),
    }

    for p in doc.paragraphs:
        for key, val in placeholders.items():
            if key in p.text:
                p.text = p.text.replace(key, val)

    word_path = "contratto_filled.docx"
    pdf_path = "contratto_filled.pdf"
    doc.save(word_path)

    convert_to_pdf(word_path, pdf_path)

    with open(pdf_path, "rb") as f:
        st.download_button("📄 Scarica PDF", f, file_name="contratto_procacciatore.pdf")

    os.remove(word_path)
    os.remove(pdf_path)
